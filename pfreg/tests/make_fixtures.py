"""Build two formula-driven V2 workbooks with the plugin's labels:
  pf2_correct.xlsx - TU from the quadratic (valid whichever constraint binds), lambda = D/TU
  pf2_buggy.xlsx   - replicates the plugin's 2026-10-06 defect: TU = D/LTV, lambda = LTV
Both are pure formulas (no cached values) so a real recalculation engine must evaluate them."""
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as L


def build(path, buggy=False):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    inp = [("Capex t=1 ($m)", 40), ("Capex t=2 ($m)", 60), ("CFADS ops year 1 ($m)", 14),
           ("CFADS growth rate", 0.02), ("Debt interest rate", 0.07), ("Debt tenor (years)", 12),
           ("Target sculpting DSCR", 1.35), ("Max debt / total uses (LTV cap)", 0.70),
           ("Upfront fee rate (% of debt)", 0.02), ("DSRA = X% of year-1 debt service", 0.5)]
    for i, (lab, v) in enumerate(inp, start=1):
        ws.cell(i, 1, lab); ws.cell(i, 2, v)
    C1, C2, CF1, G, R, _, DS, CAP, FEE, DSR = (f"$B${i}" for i in range(1, 11))
    cols = [L(c) for c in range(3, 15)]            # C..N  = ops years 1..12
    ws["A12"] = "Ops year k"
    for k, c in enumerate(cols, 1):
        ws[f"{c}12"] = k
        ws[f"{c}13"] = f"={CF1}*(1+{G})^({c}12-1)"           # CFADS
        ws[f"{c}14"] = f"={c}13/{DS}"                        # sculpted DS
        ws[f"{c}15"] = f"={c}14/(1+{R})^{c}12"               # PV
    ws["A13"], ws["A14"], ws["A15"] = "CFADS ($m)", "Sculpted DS ($m)", "PV sculpted DS ($m)"
    rows = {
        "SC": ("Sculpted capacity SC ($m)", "=SUM(C15:N15)"),
        "BETA": ("beta", f"={DSR}*{CF1}/({DS}*B20)"),
        "DLTV": ("D under LTV-binding (closed form)",
                 f"={CAP}*({C1}*(1+{R}*{CAP})+{C2})/(1-{CAP}*({FEE}*(1+{R}*{CAP})+B21))"),
        "D": ("Final debt D ($m)", "=MIN(B20,B22)"),
        "SCALE": ("scale", "=MIN(1,B23/B20)"),
        "FEEA": ("Upfront fee ($m) = rate x D", f"={FEE}*B23"),
        "DSRAA": ("DSRA ($m) = pct x DS1", f"={DSR}*C14*B24"),
        "U1": ("Uses t=1", f"={C1}+B25"),
        "TU": ("Total Uses ($m)", "=B23/" + CAP if buggy else
               f"=((B27+{C2}+B26)+SQRT((B27+{C2}+B26)^2+4*{R}*B23*B27))/2"),
        "LAM": ("λ = D / TU", f"={CAP}" if buggy else "=B23/B28"),
        "IDC": ("IDC ($m)", f"={R}*B29*B27"),
        "U2": ("Uses t=2", f"={C2}+B30+B26"),
        "DR1": ("Debt draw t=1 ($m)", "=B29*B27"),
        "DR2": ("Debt draw t=2 ($m)", "=B29*B31"),
        "EQ1": ("Equity t=1 ($m)", "=(1-B29)*B27"),
        "EQ2": ("Equity t=2 ($m)", "=(1-B29)*B31"),
        "BIND": ("Binding constraint", '=IF(B23<B20-0.000000001,"LTV cap","DSCR")'),
    }
    for i, (lab, f) in enumerate(rows.values(), start=20):
        ws.cell(i, 1, lab); ws.cell(i, 2, f)
    # debt schedule rows 40.. (C..N)
    ws["A40"], ws["A41"], ws["A42"], ws["A43"], ws["A44"] = \
        "Opening", "Interest", "Debt service", "Principal", "DSCR"
    for j, c in enumerate(cols):
        ws[f"{c}40"] = "=B23" if j == 0 else f"={cols[j-1]}40-{cols[j-1]}43"
        ws[f"{c}41"] = f"={c}40*{R}"
        ws[f"{c}42"] = f"={c}14*$B$24"
        ws[f"{c}43"] = f"={c}42-{c}41"
        ws[f"{c}44"] = f"={c}13/{c}42"
    # equity cash flows t=1..17 in C47..S47
    ws["A47"] = "Equity net CF"
    ec = [L(c) for c in range(3, 20)]
    for t, c in enumerate(ec, 1):
        if t == 1:
            f = "=-B34"
        elif t == 2:
            f = "=-B35"
        else:
            k = t - 2
            f = f"={CF1}*(1+{G})^({k}-1)" + (f"-{cols[k-1]}42" if k <= 12 else "")
            if k == 12:
                f += "+B26"
        ws[f"{c}47"] = f
    summ = [("Year-1 interest ($m)", "=C41"), ("Year-1 principal ($m)", "=C43"),
            ("Year-1 debt service ($m)", "=C42"), ("Minimum DSCR", "=MIN(C44:N44)"),
            ("LLCR at COD", "=SUMPRODUCT(C13:N13,C15:N15/C14:N14)/B23"),
            ("Equity IRR", "=IRR(C47:S47)")]
    for i, (lab, f) in enumerate(summ, start=50):
        ws.cell(i, 1, lab); ws.cell(i, 2, f)
    wb.save(path)


if __name__ == "__main__":
    d = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
    build(d / "pf2_correct.xlsx")
    build(d / "pf2_buggy.xlsx", buggy=True)
    print("fixtures written to", d)
