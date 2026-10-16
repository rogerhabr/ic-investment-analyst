"""Read metrics from / write inputs to a model workbook, and force a recalculation.

Location strategy: find a label cell by regex (case-insensitive, LAST match in workbook order,
so summary blocks at the bottom win over intermediate lines), take the first non-empty cell to
its right. Any metric can be pinned to an exact cell via a map file ({"debt": "Model!B42"}).
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

LABELS = {
    # ---- project finance outputs
    "sculpt_capacity": r"^sculpted (debt )?capacity",
    "debt": r"^final debt",
    "binding": r"^binding constraint",
    "total_uses": r"^total uses",
    "fee": r"^upfront fee \(",
    "idc": r"^idc\b",
    "dsra": r"^dsra \(",
    "gearing": r"^(λ|lambda)\b",
    "draw_t1": r"^debt draw t\s*=\s*1",
    "draw_t2": r"^debt draw t\s*=\s*2",
    "equity_t1": r"^equity (draw )?t\s*=\s*1",
    "equity_t2": r"^equity (draw )?t\s*=\s*2",
    "ds_y1": r"^year.?1 debt service",
    "int_y1": r"^year.?1 interest",
    "prin_y1": r"^year.?1 principal",
    "min_dscr": r"^min(imum)? dscr",
    "llcr": r"^llcr",
    "project_irr": r"^project.*irr",
    "equity_irr": r"^equity irr",
    # ---- DCF outputs
    "ev": r"^enterprise value",
    "pv_tv": r"^pv of terminal value",
    "pv_fcf": r"^sum of pv",
}
INPUTS = {
    "cap": r"^max debt",
    "fee": r"^upfront fee rate",
    "dsra": r"^dsra\s*=",
    "wacc": r"^wacc\b",
    "g": r"^terminal growth",
}


def _norm(v):
    return re.sub(r"\s+", " ", str(v)).strip().lower() if v is not None else ""


def _find(ws_iter, pattern):
    rx, hit = re.compile(pattern, re.I), None
    for ws in ws_iter:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and rx.search(_norm(c.value)):
                    for nxt in row[c.column:]:          # cells to the right
                        if nxt.value not in (None, ""):
                            hit = (ws.title, nxt.coordinate)
                            break
    return hit


def locate(path, keys, table, pins=None):
    wb = openpyxl.load_workbook(path, data_only=False)
    out = {}
    for k in keys:
        if pins and k in pins:
            s, a = pins[k].split("!")
            out[k] = (s, a)
        elif k in table:
            out[k] = _find(wb.worksheets, table[k])
    return out


def _parse(v):
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        m = re.search(r"-?\(?\d[\d,]*\.?\d*", v)
        if m and not re.search(r"[a-z]{3,}", v.lower()):
            s = m.group(0).replace(",", "")
            neg = s.startswith("(")
            return -float(s.strip("()")) if neg else float(s)
        return v.strip()
    return v


def read(path, locs):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for k, loc in locs.items():
        out[k] = None if loc is None else _parse(wb[loc[0]][loc[1]].value)
    return out


def formulas_present(path, locs):
    wb = openpyxl.load_workbook(path, data_only=False)
    return {k: (isinstance(wb[l[0]][l[1]].value, str) and wb[l[0]][l[1]].value.startswith("="))
            for k, l in locs.items() if l}


def write_inputs(src, dst, values: dict, pins=None):
    """Copy src -> dst with input cells overwritten. Returns {input: (sheet, cell)}."""
    locs = locate(src, values.keys(), INPUTS, pins)
    missing = [k for k, l in locs.items() if l is None]
    if missing:
        raise KeyError(f"input label not found: {missing} (pin them with --map)")
    wb = openpyxl.load_workbook(src)            # formulas kept; cached values dropped
    for k, (s, a) in locs.items():
        wb[s][a].value = values[k]
    wb.save(dst)
    return locs


# ---------------------------------------------------------------- recalculation backends
def recalc(path, backend="auto"):
    """Recalculate in place so cached values are fresh. Returns backend used."""
    if backend in ("auto", "excel") and sys.platform == "win32":
        try:
            import win32com.client as w32  # pywin32
            xl = w32.DispatchEx("Excel.Application")
            xl.Visible = False
            xl.DisplayAlerts = False
            wb = xl.Workbooks.Open(str(Path(path).resolve()))
            xl.CalculateFull()
            wb.Save()
            wb.Close(False)
            xl.Quit()
            return "excel"
        except Exception as e:  # noqa: BLE001
            if backend == "excel":
                raise RuntimeError(f"Excel COM recalculation failed: {e}") from e
    exe = shutil.which("soffice") or shutil.which("libreoffice") or \
        next((p for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",) if os.path.exists(p)), None)
    if backend in ("auto", "libreoffice") and exe:
        out = tempfile.mkdtemp()
        subprocess.run([exe, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, str(path)],
                       check=True, capture_output=True, timeout=180)
        shutil.move(os.path.join(out, Path(path).name), path)
        return "libreoffice"
    raise RuntimeError("No recalculation backend: install pywin32 (Excel) or LibreOffice")
